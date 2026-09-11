"""Build a DOCX catalogue of generated visualisation examples.

The catalogue is intended as a supplementary-material style artifact. It
includes every worked case-study PNG and every simulated mock-data PNG used by
the recommendation interface.
"""

from __future__ import annotations

import ast
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "visualisation_catalogue"
OUTPUT_DOCX = OUTPUT_DIR / "Visualisation_examples_catalogue.docx"
OUTPUT_PDF = OUTPUT_DIR / "Visualisation_examples_catalogue.pdf"

CASE_STUDY_FIGURE_DIR = ROOT / "case_study" / "figures"
MOCK_FIGURE_DIR = ROOT / "examples" / "figures"
CASE_STUDY_NOTES = CASE_STUDY_FIGURE_DIR / "case_study_visualisation_notes.md"
EXAMPLE_GENERATOR = ROOT / "examples" / "generate_mock_visualisation_examples.py"

PAGE_WIDTH_IN = 11.0
PAGE_HEIGHT_IN = 8.5
MARGIN_IN = 0.6
MAX_IMAGE_WIDTH_IN = PAGE_WIDTH_IN - (2 * MARGIN_IN)
MAX_IMAGE_HEIGHT_IN = 5.55

INK = RGBColor(23, 33, 43)
MUTED = RGBColor(86, 99, 111)
ACCENT = RGBColor(0, 114, 178)
INK_HEX = "#17212B"
MUTED_HEX = "#56636F"
ACCENT_HEX = "#0072B2"
GRID_HEX = "#D9DEDB"
CALLOUT_HEX = "#F2F7FB"


def set_cell_shading(paragraph: Paragraph, color: str) -> None:
    """Apply a light paragraph background using OOXML shading."""

    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    p_pr.append(shading)


def set_paragraph_border(paragraph: Paragraph, color: str = "D9DEDB") -> None:
    """Add a subtle bottom rule to a paragraph."""

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def fit_image(path: Path) -> tuple[float, float]:
    """Return width and height in inches fitted to the catalogue image box."""

    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    width = MAX_IMAGE_WIDTH_IN
    height = width / aspect
    if height > MAX_IMAGE_HEIGHT_IN:
        height = MAX_IMAGE_HEIGHT_IN
        width = height * aspect
    return width, height


def case_study_entries() -> list[dict[str, str | Path]]:
    """Extract ordered figure entries and captions from the case-study notes."""

    text = CASE_STUDY_NOTES.read_text(encoding="utf-8")
    pattern = re.compile(
        r"## Figure (?P<number>\d+)\. (?P<title>.*?)\n\n"
        r"\*\*Files:\*\* (?P<files>.*?)\n\n"
        r"\*\*Caption:\*\* (?P<caption>.*?)(?:\n\n|\Z)",
        re.S,
    )
    entries: list[dict[str, str | Path]] = []
    for match in pattern.finditer(text):
        files = [item.strip() for item in match.group("files").split(",")]
        png = next(file for file in files if file.endswith(".png"))
        entries.append(
            {
                "number": match.group("number"),
                "title": match.group("title").strip(),
                "caption": re.sub(r"\s+", " ", match.group("caption")).strip(),
                "path": CASE_STUDY_FIGURE_DIR / png,
                "source": "Worked case study, reproduced NHANES 2011-2014 data",
            }
        )
    return entries


def load_mock_recommendation_order() -> list[dict[str, str | Path]]:
    """Return simulated example entries in the generator's recommendation order."""

    tree = ast.parse(EXAMPLE_GENERATOR.read_text(encoding="utf-8"))
    recommendation_mapping: dict[str, dict[str, str]] | None = None
    example_functions: list[str] | None = None

    for node in tree.body:
        if not isinstance(node, ast.Assign):
            continue
        targets = [target.id for target in node.targets if isinstance(target, ast.Name)]
        if "RECOMMENDATION_TO_EXAMPLE" in targets:
            recommendation_mapping = ast.literal_eval(node.value)
        elif "EXAMPLE_FUNCTIONS" in targets:
            example_functions = [
                element.id for element in node.value.elts if isinstance(element, ast.Name)
            ]

    if recommendation_mapping is None or example_functions is None:
        raise RuntimeError(f"Could not read example metadata from {EXAMPLE_GENERATOR}")

    function_to_name = {
        metadata["function"]: name
        for name, metadata in recommendation_mapping.items()
    }

    entries: list[dict[str, str | Path]] = []
    for function_name in example_functions:
        recommendation = function_to_name[function_name]
        filename = function_name.removeprefix("example_") + ".png"
        entries.append(
            {
                "title": recommendation,
                "caption": (
                    f"Simulated data example illustrating the visual structure for "
                    f"the recommendation: {recommendation}."
                ),
                "path": MOCK_FIGURE_DIR / filename,
                "source": "Simulated mock data",
            }
        )
    return entries


def configure_document(document: Document) -> None:
    """Apply a compact report style."""

    section = document.sections[0]
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(6)

    for style_name, size, before, after in [
        ("Heading 1", 17, 16, 8),
        ("Heading 2", 12, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    caption = styles.add_style("Figure Caption", 1)
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = INK
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(2)
    caption.paragraph_format.line_spacing = 1.05

    source = styles.add_style("Figure Source", 1)
    source.font.name = "Calibri"
    source.font.size = Pt(7.5)
    source.font.italic = True
    source.font.color.rgb = MUTED
    source.paragraph_format.space_after = Pt(4)


def add_cover(document: Document, total_case: int, total_mock: int) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Visualisation Examples Catalogue")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Decision-support toolkit for visualising accelerometer-derived movement behaviour data"
    )
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        f"This document contains {total_case + total_mock} generated visualisation examples: "
        f"{total_case} worked case-study examples and {total_mock} simulated data examples."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = INK

    callout = document.add_paragraph()
    set_cell_shading(callout, "F2F7FB")
    callout.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = callout.add_run(
        "Worked case-study examples were generated from the reproduced NHANES 2011-2014 "
        "case-study outputs. Simulated data examples were generated from small mock datasets "
        "to illustrate recommendation-specific visual structures and should not be interpreted "
        "as evidence that the designs have been empirically tested."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = INK

    document.add_page_break()


def add_section_heading(document: Document, heading: str, description: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.add_run(heading)
    set_paragraph_border(paragraph)

    desc = document.add_paragraph()
    run = desc.add_run(description)
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


def add_figure_page(
    document: Document,
    title: str,
    caption: str,
    source: str,
    image_path: Path,
    figure_label: str,
) -> None:
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    heading = document.add_paragraph(style="Heading 2")
    heading.add_run(f"{figure_label}. {title}")

    width, height = fit_image(image_path)
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))

    caption_paragraph = document.add_paragraph(style="Figure Caption")
    caption_run = caption_paragraph.add_run(caption)
    caption_run.bold = False

    source_paragraph = document.add_paragraph(style="Figure Source")
    source_paragraph.add_run(f"Source: {source}.")


def build_catalogue() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    case_entries = case_study_entries()
    mock_entries = load_mock_recommendation_order()

    document = Document()
    configure_document(document)
    add_cover(document, len(case_entries), len(mock_entries))

    add_section_heading(
        document,
        "Worked case study examples",
        "Figures generated from the reproduced NHANES case-study dataset and associated summary outputs.",
    )
    for index, entry in enumerate(case_entries, start=1):
        add_figure_page(
            document,
            title=str(entry["title"]),
            caption=str(entry["caption"]),
            source=str(entry["source"]),
            image_path=Path(entry["path"]),
            figure_label=f"Worked example {index}",
        )
        document.add_page_break()

    add_section_heading(
        document,
        "Simulated data examples",
        "Figures generated from small simulated mock datasets to illustrate recommendation-specific visual structures.",
    )
    for index, entry in enumerate(mock_entries, start=1):
        add_figure_page(
            document,
            title=str(entry["title"]),
            caption=str(entry["caption"]),
            source=str(entry["source"]),
            image_path=Path(entry["path"]),
            figure_label=f"Simulated example {index}",
        )
        if index != len(mock_entries):
            document.add_page_break()

    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def wrap_pdf_text(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    """Wrap text for ReportLab drawing."""

    words = text.split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        candidate = " ".join([*current, word])
        if stringWidth(candidate, font_name, font_size) <= max_width or not current:
            current.append(word)
        else:
            lines.append(" ".join(current))
            current = [word]
    if current:
        lines.append(" ".join(current))
    return lines


def draw_wrapped_pdf_text(
    pdf: canvas.Canvas,
    text: str,
    x: float,
    y: float,
    max_width: float,
    font_name: str = "Helvetica",
    font_size: float = 10,
    color: str = INK_HEX,
    leading: float | None = None,
) -> float:
    """Draw wrapped text and return the y position below the text."""

    if leading is None:
        leading = font_size * 1.25
    pdf.setFont(font_name, font_size)
    pdf.setFillColor(HexColor(color))
    for line in wrap_pdf_text(text, font_name, font_size, max_width):
        pdf.drawString(x, y, line)
        y -= leading
    return y


def draw_pdf_cover(pdf: canvas.Canvas, total_case: int, total_mock: int) -> None:
    pdf.setPageSize(landscape(letter))
    width, height = landscape(letter)
    pdf.setFillColor(HexColor(INK_HEX))
    pdf.setFont("Helvetica-Bold", 28)
    title = "Visualisation Examples Catalogue"
    pdf.drawCentredString(width / 2, height - 145, title)

    pdf.setFont("Helvetica", 13)
    pdf.setFillColor(HexColor(MUTED_HEX))
    subtitle = (
        "Decision-support toolkit for visualising accelerometer-derived "
        "movement behaviour data"
    )
    pdf.drawCentredString(width / 2, height - 175, subtitle)

    pdf.setFont("Helvetica", 11)
    pdf.setFillColor(HexColor(INK_HEX))
    summary = (
        f"{total_case + total_mock} generated visualisation examples: "
        f"{total_case} worked case-study examples and {total_mock} simulated data examples."
    )
    pdf.drawCentredString(width / 2, height - 215, summary)

    callout_x = 90
    callout_y = height - 330
    callout_w = width - 180
    callout_h = 88
    pdf.setFillColor(HexColor(CALLOUT_HEX))
    pdf.roundRect(callout_x, callout_y, callout_w, callout_h, 8, fill=1, stroke=0)
    text = (
        "Worked case-study examples were generated from the reproduced NHANES "
        "2011-2014 case-study outputs. Simulated data examples were generated "
        "from small mock datasets to illustrate recommendation-specific visual "
        "structures and should not be interpreted as evidence that the designs "
        "have been empirically tested."
    )
    draw_wrapped_pdf_text(
        pdf,
        text,
        callout_x + 20,
        callout_y + callout_h - 26,
        callout_w - 40,
        font_size=10.5,
    )

    pdf.showPage()


def draw_pdf_section(pdf: canvas.Canvas, heading: str, description: str) -> None:
    pdf.setPageSize(landscape(letter))
    width, height = landscape(letter)
    x = MARGIN_IN * 72
    y = height - 165
    pdf.setFont("Helvetica-Bold", 24)
    pdf.setFillColor(HexColor(ACCENT_HEX))
    pdf.drawString(x, y, heading)
    pdf.setStrokeColor(HexColor(GRID_HEX))
    pdf.setLineWidth(1)
    pdf.line(x, y - 14, width - x, y - 14)
    draw_wrapped_pdf_text(
        pdf,
        description,
        x,
        y - 44,
        width - (2 * x),
        font_size=12,
        color=MUTED_HEX,
    )
    pdf.showPage()


def draw_pdf_figure_page(
    pdf: canvas.Canvas,
    title: str,
    caption: str,
    source: str,
    image_path: Path,
    figure_label: str,
) -> None:
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    aspect = image_width / image_height

    if aspect < 0.9:
        page_size = letter
        max_img_height = 570
    else:
        page_size = landscape(letter)
        max_img_height = 392

    pdf.setPageSize(page_size)
    width, height = page_size
    margin = MARGIN_IN * 72
    content_width = width - (2 * margin)

    pdf.setFont("Helvetica-Bold", 13)
    pdf.setFillColor(HexColor(ACCENT_HEX))
    title_text = f"{figure_label}. {title}"
    y = height - margin - 10
    y = draw_wrapped_pdf_text(
        pdf,
        title_text,
        margin,
        y,
        content_width,
        font_name="Helvetica-Bold",
        font_size=13,
        color=ACCENT_HEX,
        leading=15,
    )

    max_img_width = content_width
    draw_width = max_img_width
    draw_height = draw_width / aspect
    if draw_height > max_img_height:
        draw_height = max_img_height
        draw_width = draw_height * aspect

    img_x = (width - draw_width) / 2
    img_y = y - draw_height - 8
    pdf.drawImage(
        ImageReader(str(image_path)),
        img_x,
        img_y,
        width=draw_width,
        height=draw_height,
        preserveAspectRatio=True,
        mask="auto",
    )

    caption_y = img_y - 18
    caption_y = draw_wrapped_pdf_text(
        pdf,
        caption,
        margin,
        caption_y,
        content_width,
        font_size=8.7,
        color=INK_HEX,
        leading=10.5,
    )
    draw_wrapped_pdf_text(
        pdf,
        f"Source: {source}.",
        margin,
        caption_y - 3,
        content_width,
        font_name="Helvetica-Oblique",
        font_size=7.8,
        color=MUTED_HEX,
        leading=9.5,
    )
    pdf.showPage()


def build_pdf(case_entries: list[dict[str, str | Path]], mock_entries: list[dict[str, str | Path]]) -> Path:
    """Build a direct PDF version of the same catalogue."""

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(OUTPUT_PDF), pagesize=landscape(letter))
    pdf.setTitle("Visualisation Examples Catalogue")
    pdf.setAuthor("Visualisation decision-support toolkit")
    pdf.setSubject("Generated worked case-study and simulated visualisation examples")

    draw_pdf_cover(pdf, len(case_entries), len(mock_entries))
    draw_pdf_section(
        pdf,
        "Worked case study examples",
        "Figures generated from the reproduced NHANES case-study dataset and associated summary outputs.",
    )
    for index, entry in enumerate(case_entries, start=1):
        draw_pdf_figure_page(
            pdf,
            title=str(entry["title"]),
            caption=str(entry["caption"]),
            source=str(entry["source"]),
            image_path=Path(entry["path"]),
            figure_label=f"Worked example {index}",
        )

    draw_pdf_section(
        pdf,
        "Simulated data examples",
        "Figures generated from small simulated mock datasets to illustrate recommendation-specific visual structures.",
    )
    for index, entry in enumerate(mock_entries, start=1):
        draw_pdf_figure_page(
            pdf,
            title=str(entry["title"]),
            caption=str(entry["caption"]),
            source=str(entry["source"]),
            image_path=Path(entry["path"]),
            figure_label=f"Simulated example {index}",
        )

    pdf.save()
    return OUTPUT_PDF


def build_all() -> tuple[Path, Path]:
    """Build both DOCX and PDF catalogue artifacts."""

    docx_path = build_catalogue()
    pdf_path = build_pdf(case_study_entries(), load_mock_recommendation_order())
    return docx_path, pdf_path


if __name__ == "__main__":
    for path in build_all():
        print(path)
